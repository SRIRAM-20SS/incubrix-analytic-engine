"""
IncuBrix Citation Engine: Exporter & Workbook Synchronizer (Phase 5)
Exports the 30-target dataset to CSV, generates the comprehensive technical report,
and populates rows 15 to 44 (items C01 to C30) in submission workbook tab 02c_PAGES_CITATIONS.
"""

import os
from typing import List, Dict, Any
import pandas as pd
import openpyxl

from src.citations.discover_targets import TargetDiscoverer
from src.citations.prioritize_targets import TargetPrioritizer
from src.citations.analyze_opportunities import OpportunityAnalyzer


def export_citation_targets_csv(output_path: str = "outputs/citations/citation_targets.csv") -> str:
    """Exports all 30 citation targets to CSV with all 21 assessment fields."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    targets = TargetPrioritizer.rank_all_targets()

    # Format related_question_ids as comma-separated string for CSV
    clean_targets = []
    for t in targets:
        row = dict(t)
        if isinstance(row.get("related_question_ids"), list):
            row["related_question_ids"] = ", ".join(row["related_question_ids"])
        clean_targets.append(row)

    # Sort sequentially by target_id (C01 to C30) for clean CSV records
    clean_targets.sort(key=lambda x: int(x["target_id"].replace("C", "")))

    columns = [
        "target_id",
        "publisher_name",
        "page_title",
        "page_url",
        "page_type",
        "related_question_ids",
        "target_question",
        "market",
        "discovery_method",
        "observed_in_google",
        "observed_in_ai_answer",
        "citation_evidence",
        "competing_brands_mentioned",
        "page_relevance",
        "editorial_context",
        "potential_inclusion_angle",
        "suggested_ethical_action",
        "priority",
        "date_checked",
        "evidence_reference",
        "verification_status"
    ]

    df = pd.DataFrame(clean_targets)[columns]
    df.to_csv(output_path, index=False)
    return output_path


def generate_citation_report(output_md: str = "outputs/citations/citation_opportunity_report.md") -> str:
    """Generates the comprehensive Phase 5 Markdown opportunity report."""
    os.makedirs(os.path.dirname(output_md), exist_ok=True)
    ranked_targets = TargetPrioritizer.rank_all_targets()
    stats = OpportunityAnalyzer.get_summary_statistics()

    md = [
        "# IncuBrix Phase 5: Third-Party Citation Opportunity Research",
        "",
        "**Assessment Component**: Option 3 — AEO, SEO & GEO Visibility Engine  ",
        "**Date Checked**: `2026-10-24`  ",
        f"**Total Documented Targets**: Exactly {len(ranked_targets)} Verified Third-Party Opportunities  ",
        "**Discovery Standard**: Empirical observations from Phase 2 Google Search baseline, Perplexity citations, and creator authorities  ",
        "**Integrity Standard**: 100% white-hat PR and educational outreach. Zero backlink purchasing, zero automated spam, zero fabricated reviews.",
        "",
        "---",
        "",
        "## 1. Executive Summary & Discovery Intelligence",
        "",
        f"Phase 5 identified **{len(ranked_targets)} distinct, high-credibility third-party digital destinations** that search engines and AI assistants actively query and cite when answering the 100 creator questions cataloged in Phase 1.",
        "",
        "### Key Discovery Insights:",
        f"- **AI Assistant Citations**: **{stats['observed_in_ai_assistant']} of {stats['total_targets']} targets** were directly observed as citations in Perplexity Sonar, ChatGPT Search, Claude 3.5 Sonnet, or Google AI Overviews during Phase 2 baseline runs.",
        f"- **Search Engine Visibility**: **{stats['observed_in_google_search']} of {stats['total_targets']} targets** occupy Top 5 organic rankings or Google Discussions SERP modules for core creator workflow queries.",
        "- **IncuBrix Baseline Presence**: As established in Phase 2, IncuBrix is currently **unmentioned across 100% of these 30 targets**. Competitors (OpusClip, Descript, Submagic, Kapwing, Riverside) heavily dominate these pages.",
        "- **The Strategic Imperative**: Rather than purchasing low-value backlinks, IncuBrix must execute an ethical, high-utility digital PR and product documentation strategy to earn legitimate mentions, category listings, and citations in these high-authority hubs.",
        "",
        "---",
        "",
        "## 2. Priority Ranking & Target Classification",
        "",
        "Targets are ranked using our documented 5-factor scoring framework (Relevance, Credibility, Discoverability, Competitor Density, and Ethical Action Feasibility):",
        "",
        "| Target ID | Priority | Score | Publisher Name | Page Title | Format | Discovery Source |",
        "| :---: | :---: | :---: | :--- | :--- | :--- | :--- |"
    ]

    for t in ranked_targets:
        md.append(
            f"| **{t['target_id']}** | `{t['priority']}` | `{t['composite_score']}/25` | "
            f"{t['publisher_name']} | [{t['page_title'][:40]}...]({t['page_url']}) | "
            f"{t['page_type']} | {t['discovery_method'][:28]}... |"
        )

    md.extend([
        "",
        "---",
        "",
        "## 3. Detailed Target Analyses (All 30 Opportunities)",
        ""
    ])

    for t in ranked_targets:
        analysis = OpportunityAnalyzer.analyze_single_target(t)
        qids = ", ".join(t["related_question_ids"]) if isinstance(t["related_question_ids"], list) else t["related_question_ids"]
        md.extend([
            f"### {t['target_id']}: {t['publisher_name']} — *{t['page_title']}*",
            "",
            f"- **Target URL**: [{t['page_url']}]({t['page_url']})",
            f"- **Priority**: **{t['priority']}** (Composite Score: `{t['composite_score']}/25`) | **Market**: `{t['market']}`",
            f"- **Page Type**: `{t['page_type']}` | **Related Question IDs**: `{qids}`",
            f"- **Target Question**: *\"{t['target_question']}\"*",
            f"- **Discovery Evidence**: {analysis['2_observation_evidence']}",
            f"- **Current Competing Brands Present**: `{t['competing_brands_mentioned']}`",
            f"- **Current IncuBrix Presence**: `{analysis['4_incubrix_current_presence']}`",
            "",
            "**Strategic Opportunity Analysis:**",
            f"1. **Relevance Explanation**: {analysis['1_relevance_explanation']}",
            f"2. **Current Content Utility**: {analysis['3_current_content_utility']}",
            f"3. **Factual Contribution Angle**: {analysis['6_factual_contribution_angle']}",
            f"4. **Suggested Ethical Action**: {analysis['7_suggested_ethical_action']}",
            "",
            "---",
            ""
        ])

    md.extend([
        "## 4. Ethical PR & Inclusion Action Guidelines",
        "",
        "IncuBrix adheres to strict academic and commercial integrity standards:",
        "1. **Zero Link Schemes**: IncuBrix will never engage in paid link networks, sponsored PBNs, or grey-hat link exchanges.",
        "2. **Directory Claiming**: Directory targets (`C11` G2, `C12` Capterra, `C14` Toolify) should be claimed legitimately via verified business credentials with transparent pricing and credit specifications.",
        "3. **Editorial Contribution**: Tech publications (`C01` Zapier, `C04` Influencer Marketing Hub, `C10` TechRadar, `C16` PCMag) should be pitched only with original data benchmarks, technical case studies (e.g. saving $595/mo per client for agencies), or verified product documentation.",
        "4. **Community Transparency**: Community forums (`C27` Reddit r/VideoEditing, `C28` Reddit r/ContentCreation) should be engaged solely by identified technical staff answering encoding questions (VFR vs CFR, safe-zone layouts) without surreptitious astroturfing.",
        "",
        "---",
        "",
        "## 5. Instructions for Adding & Verifying New Targets",
        "",
        "To add new citation targets to the IncuBrix Visibility Engine:",
        "1. Open `src/citations/discover_targets.py` or import a CSV into `src/citations/validate_targets.py`.",
        "2. Ensure all 21 mandatory fields are populated, including real URLs, question mappings, and discovery evidence.",
        "3. Run the automated validator `python -m src.citations.validate_targets` to verify zero duplicate IDs or URLs.",
        "4. Execute `python run_citations.py` to re-rank opportunities, regenerate the report, and update the workbook tab `02c_PAGES_CITATIONS`.",
        "",
        "---",
        "",
        "## 6. Replication Commands",
        "",
        "```bash",
        "# 1. Run the end-to-end Phase 5 citation discovery and ranking engine",
        "python run_citations.py",
        "",
        "# 2. Run the automated Phase 5 verification test suite",
        "pytest -v tests/test_citations.py",
        "",
        "# 3. Run the full test suite across all 5 phases",
        "pytest -v",
        "```",
        ""
    ])

    with open(output_md, "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    return output_md


def populate_workbook_tab_02c_citations(
    workbook_path: str = "submission/IncuBrix_2027_Growth_Engineer_Option_3_AEO_SEO_GEO_Submission_Workbook.xlsx"
) -> str:
    """Populates rows 15 to 44 (items C01 to C30) in tab 02c_PAGES_CITATIONS."""
    if not os.path.exists(workbook_path):
        return ""

    wb = openpyxl.load_workbook(workbook_path)
    if "02c_PAGES_CITATIONS" not in wb.sheetnames:
        return ""

    sheet = wb["02c_PAGES_CITATIONS"]
    targets = TargetPrioritizer.rank_all_targets()
    # Sort sequentially by target_id (C01 to C30) for clean workbook navigation
    targets.sort(key=lambda x: int(x["target_id"].replace("C", "")))

    # Rows 15 to 44 correspond to C01 to C30
    for idx, t in enumerate(targets):
        row = 15 + idx
        qids = ", ".join(t["related_question_ids"]) if isinstance(t["related_question_ids"], list) else t["related_question_ids"]

        sheet.cell(row, 1, t["target_id"])
        sheet.cell(row, 2, "Citation target")
        sheet.cell(row, 3, qids)
        sheet.cell(row, 4, t["page_title"])
        sheet.cell(row, 5, t["page_url"])
        sheet.cell(row, 6, f"{t['page_relevance']} ({t['discovery_method']})")
        sheet.cell(row, 7, t["potential_inclusion_angle"])
        sheet.cell(row, 8, f"Publisher: {t['publisher_name']} | Competitors: {t['competing_brands_mentioned']}")
        sheet.cell(row, 9, "Ready")

    wb.save(workbook_path)
    return workbook_path


if __name__ == "__main__":
    p_csv = export_citation_targets_csv()
    p_md = generate_citation_report()
    p_wb = populate_workbook_tab_02c_citations()
    print(f"Exported citation targets CSV to: {p_csv}")
    print(f"Generated citation opportunity report at: {p_md}")
    if p_wb:
        print(f"Updated workbook at: {p_wb}")
