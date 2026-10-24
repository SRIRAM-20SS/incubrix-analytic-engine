"""
Automated Test Suite for Phase 5: Third-Party Citation Opportunity Research
Tests target count (exactly 30), ID formatting, URL validity, schema completeness,
deduplication, prioritization scoring, CSV export, and workbook tab 02c_PAGES_CITATIONS.
"""

import os
import re
import pytest
import pandas as pd
import openpyxl

from src.citations.discover_targets import TargetDiscoverer, CITATION_TARGETS_RAW
from src.citations.validate_targets import TargetValidator, REQUIRED_FIELDS
from src.citations.prioritize_targets import TargetPrioritizer
from src.citations.analyze_opportunities import OpportunityAnalyzer
from src.citations.export_citations import export_citation_targets_csv, generate_citation_report, populate_workbook_tab_02c_citations


class TestCitationDatasetIntegrity:
    def test_exactly_thirty_targets_collected(self):
        targets = TargetDiscoverer.get_all_targets()
        assert len(targets) == 30, f"Expected 30 targets, found {len(targets)}"

    def test_all_target_ids_unique_and_formatted(self):
        targets = TargetDiscoverer.get_all_targets()
        ids = [t["target_id"] for t in targets]
        assert len(set(ids)) == 30, "Duplicate target IDs found"
        for i in range(1, 31):
            expected = f"C{i:02d}"
            assert expected in ids, f"Missing {expected} in target IDs"

    def test_all_target_urls_unique_and_valid(self):
        targets = TargetDiscoverer.get_all_targets()
        urls = [t["page_url"] for t in targets]
        assert len(set(urls)) == 30, f"Duplicate target URLs detected: {set(urls)}"
        for u in urls:
            assert TargetValidator.validate_url(u) is True, f"Invalid URL: {u}"
            assert u.startswith("https://") or u.startswith("http://")

    def test_all_twenty_one_fields_present_and_non_empty(self):
        targets = TargetDiscoverer.get_all_targets()
        for t in targets:
            for field in REQUIRED_FIELDS:
                assert field in t, f"Target {t['target_id']} missing field '{field}'"
                val = t[field]
                if isinstance(val, list):
                    assert len(val) > 0, f"Target {t['target_id']} field '{field}' is empty list"
                else:
                    assert str(val).strip() != "", f"Target {t['target_id']} field '{field}' is empty"

    def test_dataset_passes_target_validator(self):
        targets = TargetDiscoverer.get_all_targets()
        res = TargetValidator.validate_dataset(targets)
        assert res["is_valid"] is True, f"Validation errors: {res['errors']}"
        assert res["total_targets"] == 30
        assert res["unique_ids_count"] == 30
        assert res["unique_urls_count"] == 30


class TestPrioritizationAndAnalysis:
    def test_prioritizer_calculates_valid_scores(self):
        ranked = TargetPrioritizer.rank_all_targets()
        assert len(ranked) == 30
        for r in ranked:
            assert "composite_score" in r
            assert 5 <= r["composite_score"] <= 25
            assert r["priority"] in ("Critical", "High", "Medium", "Low")

    def test_opportunity_analyzer_seven_dimensions(self):
        analyses = OpportunityAnalyzer.analyze_all_targets()
        assert len(analyses) == 30
        required_analysis_keys = [
            "1_relevance_explanation",
            "2_observation_evidence",
            "3_current_content_utility",
            "4_incubrix_current_presence",
            "5_competing_products_present",
            "6_factual_contribution_angle",
            "7_suggested_ethical_action"
        ]
        for a in analyses:
            for k in required_analysis_keys:
                assert k in a, f"Analysis for {a['target_id']} missing {k}"
                assert len(str(a[k]).strip()) > 10

    def test_ethical_action_does_not_contain_prohibited_terms(self):
        targets = TargetDiscoverer.get_all_targets()
        forbidden = ["buy link", "paid backlink", "pbn", "fake review", "automated spam", "bot comment"]
        for t in targets:
            action = t["suggested_ethical_action"].lower()
            for f in forbidden:
                assert f not in action, f"Target {t['target_id']} has forbidden ethical action: '{f}'"


class TestExportDeliverables:
    def test_csv_export_has_thirty_rows_and_twenty_one_columns(self):
        csv_path = export_citation_targets_csv("outputs/citations/citation_targets.csv")
        assert os.path.exists(csv_path)
        df = pd.read_csv(csv_path)
        assert len(df) == 30
        for f in REQUIRED_FIELDS:
            assert f in df.columns, f"CSV missing required column '{f}'"

    def test_markdown_report_generation(self):
        md_path = generate_citation_report("outputs/citations/citation_opportunity_report.md")
        assert os.path.exists(md_path)
        with open(md_path, "r", encoding="utf-8") as f:
            content = f.read()
        for i in range(1, 31):
            assert f"C{i:02d}" in content, f"Report missing target C{i:02d}"

    def test_workbook_tab_02c_populated_with_thirty_targets(self):
        wb_path = populate_workbook_tab_02c_citations("submission/IncuBrix_2027_Growth_Engineer_Option_3_AEO_SEO_GEO_Submission_Workbook.xlsx")
        assert os.path.exists(wb_path)
        wb = openpyxl.load_workbook(wb_path, data_only=True)
        assert "02c_PAGES_CITATIONS" in wb.sheetnames
        sheet = wb["02c_PAGES_CITATIONS"]
        # Rows 15 to 44 correspond to C01 to C30
        for idx in range(30):
            row = 15 + idx
            expected_id = f"C{idx+1:02d}"
            assert sheet.cell(row, 1).value == expected_id
            assert sheet.cell(row, 2).value == "Citation Target"
            assert sheet.cell(row, 9).value == "Target Identified"
