"""
Automated Test Suite for Phase 3: Technical SEO Audit & 20 Validated Fixes
Tests fix count, schema completeness, code artifact existence, validator execution,
XML and JSON-LD parsing, and workbook persistence.
"""

import os
import pytest
import openpyxl
import pandas as pd
import json
import xml.etree.ElementTree as ET

from src.seo_audit.fix_generator import get_all_technical_fixes, export_fixes_csv, generate_technical_audit_report, populate_workbook_fixes
from src.seo_audit.validators import FixValidator
from src.seo_audit.page_metadata import MetadataAuditor
from src.seo_audit.robots_sitemap import RobotsSitemapAuditor
from src.seo_audit.performance import PerformanceAuditor
from src.seo_audit.structured_data import SchemaAuditor


class TestTechnicalFixesDataset:
    @pytest.fixture(scope="module")
    def fixes_list(self):
        return get_all_technical_fixes()

    def test_exactly_twenty_fixes(self, fixes_list):
        assert len(fixes_list) == 20, f"Expected exactly 20 technical fixes, found {len(fixes_list)}"

    def test_all_fix_ids_unique_and_formatted(self, fixes_list):
        ids = [f["fix_id"] for f in fixes_list]
        assert len(set(ids)) == 20, "Duplicate fix IDs detected"
        for i in range(1, 21):
            expected_id = f"FIX-{i:02d}"
            assert expected_id in ids, f"Expected {expected_id} in fixes dataset"

    def test_mandatory_fields_present_and_non_empty(self, fixes_list):
        required_fields = [
            "fix_id",
            "affected_url",
            "audit_category",
            "observed_problem",
            "evidence",
            "date_checked",
            "impact",
            "priority",
            "recommended_solution",
            "file_artifact",
            "validation_method",
            "validation_result",
            "implementation_status",
            "limitations"
        ]
        for f in fixes_list:
            for field in required_fields:
                assert field in f, f"Fix {f['fix_id']} missing field '{field}'"
                assert str(f[field]).strip() != "", f"Fix {f['fix_id']} field '{field}' is empty"

    def test_code_artifacts_exist_and_non_empty(self, fixes_list):
        for f in fixes_list:
            path = f["file_artifact"]
            assert os.path.exists(path), f"Code artifact for {f['fix_id']} not found at {path}"
            assert os.path.getsize(path) > 50, f"Code artifact for {f['fix_id']} is too small or empty"

    def test_implementation_code_read_successfully(self, fixes_list):
        for f in fixes_list:
            code = f.get("implementation_code", "")
            assert len(code.strip()) > 30, f"Fix {f['fix_id']} has empty implementation_code"


class TestFixValidators:
    @pytest.fixture(scope="module")
    def fixes_list(self):
        return get_all_technical_fixes()

    def test_all_twenty_artifacts_pass_validation(self, fixes_list):
        for f in fixes_list:
            code = f.get("implementation_code", "")
            fn = os.path.basename(f["file_artifact"])
            res = FixValidator.validate_artifact(f["fix_id"], code, fn)
            assert res["is_valid"] is True, f"Fix {f['fix_id']} failed validation: {res['message']}"
            assert res["errors"] == 0, f"Fix {f['fix_id']} has {res['errors']} validation errors"

    def test_json_ld_schema_tags_parse_valid_json(self, fixes_list):
        json_ld_fixes = ["FIX-01", "FIX-03", "FIX-08", "FIX-09", "FIX-12"]
        for fix_id in json_ld_fixes:
            fix = next(f for f in fixes_list if f["fix_id"] == fix_id)
            code = fix["implementation_code"]
            # Extract script content
            start = code.find("<script type=\"application/ld+json\">") + len("<script type=\"application/ld+json\">")
            end = code.find("</script>")
            assert start != -1 and end != -1, f"Fix {fix_id} missing <script type=\"application/ld+json\"> tags"
            json_text = code[start:end].strip()
            data = json.loads(json_text)
            assert isinstance(data, dict), f"Fix {fix_id} JSON-LD is not an object"
            assert "@context" in data or "@graph" in data, f"Fix {fix_id} missing @context"

    def test_xml_sitemap_is_well_formed(self, fixes_list):
        fix_04 = next(f for f in fixes_list if f["fix_id"] == "FIX-04")
        code = fix_04["implementation_code"]
        root = ET.fromstring(code.strip())
        assert root.tag.endswith("urlset"), f"Expected <urlset> root tag, got {root.tag}"
        urls = root.findall("{http://www.sitemaps.org/schemas/sitemap/0.9}url")
        assert len(urls) >= 7, f"Expected at least 7 URLs in sitemap, got {len(urls)}"

    def test_llms_txt_structure(self, fixes_list):
        fix_02 = next(f for f in fixes_list if f["fix_id"] == "FIX-02")
        lines = fix_02["implementation_code"].splitlines()
        assert any(l.startswith("# IncuBrix") for l in lines)
        assert any("Core Capabilities" in l for l in lines)


class TestExportDeliverables:
    def test_fixes_csv_export(self):
        path = export_fixes_csv("outputs/fixes/fixes.csv")
        assert os.path.exists(path)
        df = pd.read_csv(path)
        assert len(df) == 20
        assert "fix_id" in df.columns
        assert "observed_problem" in df.columns
        assert "validation_method" in df.columns

    def test_technical_audit_md_report(self):
        path = generate_technical_audit_report("outputs/fixes/technical_audit.md")
        assert os.path.exists(path)
        with open(path, "r", encoding="utf-8") as f:
            content = f.read()
        for i in range(1, 21):
            assert f"FIX-{i:02d}" in content, f"Report missing FIX-{i:02d}"

    def test_workbook_tab_02b_populated(self):
        wb_path = populate_workbook_fixes("submission/IncuBrix_2027_Growth_Engineer_Option_3_AEO_SEO_GEO_Submission_Workbook.xlsx")
        assert os.path.exists(wb_path)
        wb = openpyxl.load_workbook(wb_path, data_only=True)
        assert "02b_FIXES" in wb.sheetnames
        sheet = wb["02b_FIXES"]
        # Check first and 20th fix row
        assert sheet.cell(5, 1).value == "FIX-01"
        assert sheet.cell(24, 1).value == "FIX-20"
        assert sheet.cell(5, 8).value == "Yes"


class TestAuditorEngines:
    def test_metadata_auditor_detects_empty_tags(self):
        bad_html = "<html><head><title>Short</title></head><body><h1>Heading</h1></body></html>"
        res = MetadataAuditor.audit_html(bad_html, "https://test.com")
        findings = [f["type"] for f in res["findings"]]
        assert "TitleTooShort" in findings
        assert "MetaDescriptionMissing" in findings

    def test_robots_auditor_detects_missing_ai_crawlers(self):
        robots_text = "User-agent: *\nAllow: /\nSitemap: https://incubrix.com/sitemap.xml"
        res = RobotsSitemapAuditor.audit_robots_txt(robots_text)
        assert res["has_sitemap"] is True
        assert res["has_ai_crawlers"] is False
        findings = [f["type"] for f in res["findings"]]
        assert "AICrawlersUnspecified" in findings

    def test_schema_auditor_flags_missing_schemas(self):
        empty_html = "<html><head><title>FAQ</title></head><body><h1>FAQ</h1></body></html>"
        res = SchemaAuditor.audit_page_schemas(empty_html, "https://incubrix.com/faq")
        assert res["json_ld_blocks_count"] == 0
        findings = [f["type"] for f in res["findings"]]
        assert "NoStructuredData" in findings
